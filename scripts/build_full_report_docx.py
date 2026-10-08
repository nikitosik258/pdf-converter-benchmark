"""Create a self-contained Word edition of the verified Prompt 15 report.

The input is the already published Markdown report and its local assets.  This
is deliberately a presentation-only operation: it does not read or alter PDF
files, Ground Truth, benchmark caches, adapter outputs, or scores.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def add_hyperlink(paragraph, label: str, url: str) -> None:
    """Append a clickable external or relative hyperlink to a paragraph."""
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_markdown_text(paragraph, text: str) -> None:
    """Write ordinary Markdown inline text, retaining all links as hyperlinks."""
    position = 0
    for match in LINK_RE.finditer(text):
        paragraph.add_run(text[position:match.start()])
        add_hyperlink(paragraph, match.group(1), match.group(2))
        position = match.end()
    tail = text[position:]
    # Captions use one pair of emphasis markers only in this source report.
    if tail.startswith("*") and tail.endswith("*") and len(tail) > 2:
        paragraph.add_run(tail[1:-1]).italic = True
    else:
        paragraph.add_run(tail)


def split_table_row(line: str) -> list[str]:
    cells = re.split(r"(?<!\\)\|", line.strip())
    if cells and not cells[0].strip():
        cells = cells[1:]
    if cells and not cells[-1].strip():
        cells = cells[:-1]
    return [cell.strip().replace("\\|", "|").replace("<br>", "\n") for cell in cells]


def is_separator(cells: list[str]) -> bool:
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def add_table(document: Document, lines: list[str]) -> None:
    rows = [split_table_row(line) for line in lines]
    if len(rows) < 2 or not is_separator(rows[1]):
        for line in lines:
            document.add_paragraph(line)
        return
    headers, body = rows[0], rows[2:]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.size = Pt(7.5)
    for source in body:
        row = table.add_row()
        for index in range(len(headers)):
            cell = row.cells[index]
            cell.text = ""
            value = source[index] if index < len(source) else ""
            paragraph = cell.paragraphs[0]
            add_markdown_text(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(7.5)


def add_picture(document: Document, report_dir: Path, markdown_path: str, caption: str) -> None:
    picture = report_dir / markdown_path
    if picture.exists():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(picture), width=Cm(23.5))
    else:
        document.add_paragraph(f"[Не найдено изображение: {markdown_path}]")
    if caption:
        note = document.add_paragraph(style="Caption")
        note.alignment = WD_ALIGN_PARAGRAPH.CENTER
        note.add_run(caption)


def configure_document(document: Document, experiment_id: str) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10)
    for name, size, color in (("Title", 22, "15394A"), ("Heading 1", 16, "15394A"), ("Heading 2", 13, "15394A")):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    if "Formula" not in [style.name for style in document.styles]:
        formula = document.styles.add_style("Formula", WD_STYLE_TYPE.PARAGRAPH)
        formula.font.name = "Cambria Math"
        formula.font.size = Pt(11)
        formula.paragraph_format.space_before = Pt(6)
        formula.paragraph_format.space_after = Pt(6)
    header = section.header.paragraphs[0]
    header.text = "PDF Converter Benchmark · полный технический отчёт"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.text = f"Baseline {experiment_id} · Prompt 15 · Word edition"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.core_properties.title = "Сравнительное исследование средств анализа PDF и преобразования в текст"
    document.core_properties.subject = f"Воспроизводимый Word-отчёт для benchmark {experiment_id}"
    document.core_properties.author = "PDF Converter Benchmark"


def render_markdown(markdown: str, document: Document, report_dir: Path) -> None:
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(line[2:].strip())
            index += 1
            continue
        if line.startswith("## "):
            document.add_page_break()
            document.add_heading(line[3:].strip(), level=1)
            index += 1
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=2)
            index += 1
            continue
        image = IMAGE_RE.fullmatch(line.strip())
        if image:
            add_picture(document, report_dir, image.group(2), image.group(1))
            index += 1
            continue
        if line == "$$":
            formula_lines: list[str] = []
            index += 1
            while index < len(lines) and lines[index] != "$$":
                formula_lines.append(lines[index])
                index += 1
            formula = document.add_paragraph(style="Formula")
            formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
            formula.add_run("\n".join(formula_lines))
            index += 1
            continue
        if line.startswith("```"):
            code_lines: list[str] = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index])
                index += 1
            paragraph = document.add_paragraph()
            run = paragraph.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            index += 1
            continue
        if line.startswith("|"):
            table_lines: list[str] = []
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
            continue
        if line.startswith("- "):
            while index < len(lines) and lines[index].startswith("- "):
                paragraph = document.add_paragraph(style="List Bullet")
                add_markdown_text(paragraph, lines[index][2:])
                index += 1
            continue
        paragraph_lines = [line]
        index += 1
        while index < len(lines) and lines[index].strip() and not any((lines[index].startswith(prefix) for prefix in ("# ", "## ", "### ", "- ", "|", "$$", "```", "!["))):
            paragraph_lines.append(lines[index])
            index += 1
        paragraph = document.add_paragraph()
        add_markdown_text(paragraph, " ".join(part.strip() for part in paragraph_lines))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--experiment-id", required=True)
    args = parser.parse_args()
    if not re.fullmatch(r"[A-Za-z0-9_-]+", args.experiment_id):
        raise ValueError("Invalid experiment ID")
    report_dir = ROOT / "reports" / "final_report" / args.experiment_id
    source = report_dir / "report.md"
    destination = report_dir / "full_technical_report.docx"
    if not source.exists():
        raise FileNotFoundError(source)
    document = Document()
    configure_document(document, args.experiment_id)
    render_markdown(source.read_text(encoding="utf-8"), document, report_dir)
    document.save(destination)
    headings = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    if len(headings) != 20:
        raise RuntimeError(f"Expected 20 report sections, got {len(headings)}")
    if destination.stat().st_size < 100_000:
        raise RuntimeError("DOCX unexpectedly small")
    print(destination)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
