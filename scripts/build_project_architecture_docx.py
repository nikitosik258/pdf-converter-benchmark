"""Create a short Word description of the PDF benchmark architecture.

This presentation-only script documents the project and does not run adapters,
access cloud APIs, or change benchmark inputs/results.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "PROJECT_ARCHITECTURE.docx"


def set_font(run, size: float | None = None, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size:
        run.font.size = Pt(size)
    run.bold = bold


def paragraph(document: Document, text: str = "", *, bold: bool = False) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    run = p.add_run(text)
    set_font(run, 11, bold)


def bullet(document: Document, text: str) -> None:
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_font(run, 10.5)


def heading(document: Document, text: str, level: int = 1) -> None:
    p = document.add_heading(level=level)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_font(run, 15 if level == 1 else 12.5, True)
    run.font.color.rgb = RGBColor(21, 57, 74)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(value)
        set_font(run, 9.5, True)
    for source in rows:
        cells = table.add_row().cells
        for i, value in enumerate(source):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(value)
            set_font(run, 9.5)


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(11)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Краткое описание архитектуры проекта\nPDF Converter Benchmark")
    set_font(run, 18, True)
    run.font.color.rgb = RGBColor(21, 57, 74)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Исследование качества преобразования технических PDF в текстовые и структурированные данные")
    set_font(run, 11)

    heading(document, "1. Назначение проекта")
    paragraph(
        document,
        "Проект сравнивает пять облачных сервисов и пять локальных библиотек для обработки PDF. "
        "Оценивается не только обычный текст, но и таблицы, математические и химические формулы, изображения и диаграммы. "
        "Результат исследования — воспроизводимый benchmark, отчёт с количественными оценками и web-приложение для демонстрации конвертеров.",
    )

    heading(document, "2. Что и как делается в проекте")
    paragraph(document, "Работа построена как единый воспроизводимый цикл:")
    for item in [
        "Пять технических PDF (D01–D05) образуют тестовый корпус: всего 98 страниц и 40 заранее размеченных эталонных объектов.",
        "Каждый из десяти конвертеров извлекает содержимое PDF. Их ответы приводятся к общей внутренней структуре StandardizedDocument.",
        "Извлечённые объекты сопоставляются с Ground Truth по типу, тексту и координатам на странице. Для каждого объекта рассчитывается оценка качества от 0 до 100.",
        "Оценки агрегируются по шести категориям, документам и инструментам; затем строятся таблицы, графики, error analysis и полный технический отчёт.",
        "Сохранённый кэш результатов позволяет повторно анализировать и экспортировать результаты без новых платных облачных запросов и без повторного запуска тяжёлых моделей.",
    ]:
        bullet(document, item)

    heading(document, "3. Логическая архитектура")
    add_table(document, ["Слой", "Состав", "Роль"], [
        ["Исходные данные", "documents/, ground_truth/, object_manifest.json", "Тестовые PDF и эталонная разметка объектов."],
        ["Адаптеры", "src/pdf_benchmark/adapters/", "Единый интерфейс для локальных библиотек и облачных API."],
        ["Нормализация", "StandardizedDocument", "Приведение разных ответов к тексту, таблицам, формулам, изображениям и диаграммам с координатами."],
        ["Оценивание", "evaluation/, benchmark/", "Сопоставление с Ground Truth, расчёт метрик, агрегация и контроль качества."],
        ["Кэш и результаты", "outputs/benchmark/, reports/, exports/", "Хранение версионированных результатов, отчётов и TXT-экспортов."],
        ["Web-приложение", "src/pdf_benchmark/web/", "Загрузка PDF, запуск одного конвертера, просмотр и скачивание результата."],
    ])

    heading(document, "4. Поток данных benchmark")
    flow = document.add_paragraph()
    flow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = flow.add_run("PDF D01–D05  →  Adapter  →  Raw response/cache  →  StandardizedDocument\n→  Matching с Ground Truth  →  Метрики по 6 категориям  →  Таблицы, графики, отчёт")
    set_font(run, 11, True)
    flow.paragraph_format.space_before = Pt(8)
    flow.paragraph_format.space_after = Pt(8)
    paragraph(document, "Категории оценивания: текст, таблицы, математика, химия, изображения и диаграммы. Overall Score используется только как сводный показатель; выводы делаются также по отдельным категориям и документам.")

    heading(document, "5. Инструменты")
    add_table(document, ["Группа", "Инструменты"], [
        ["Локальные", "PyMuPDF, pdfplumber, pdfminer.six, Docling, MinerU"],
        ["Облачные", "OCR.Space, Nutrient, Mindee, Adobe Extract, LlamaParse"],
    ])
    paragraph(document, "Облачные адаптеры изолированы: ключи берутся только из переменных окружения и не сохраняются в исходном коде или в интерфейсе.")

    heading(document, "6. Архитектура web-приложения")
    paragraph(document, "Текущая демонстрационная реализация — однопользовательское приложение FastAPI со статическим браузерным интерфейсом.")
    add_table(document, ["Компонент", "Функция"], [
        ["Браузерный интерфейс", "Загрузка PDF, выбор доступного converter, отображение статуса, времени, текста, структурированных данных и ошибок."],
        ["FastAPI", "REST API: upload, список converters, запуск задания, статус, результат, скачивание TXT/JSON, healthcheck."],
        ["Фоновый runner", "Запускает существующий benchmark worker в отдельном процессе, не дублируя адаптеры."],
        ["Временное хранилище", "Сохраняет входной PDF, состояние задания и результаты на время работы; предусмотрена автоматическая очистка."],
        ["Docker", "CPU-режим для лёгких инструментов и GPU-режим для тяжёлых ML-конвертеров."],
    ])

    heading(document, "7. Воспроизводимость и ограничения")
    for item in [
        "Официальный baseline: 20261004T095716Z_s42_2334c932. В нём 10 инструментов, 50 пар «инструмент–PDF», 400 оцениваемых объектов и 0 структурных ошибок.",
        "Ground Truth, исходные PDF и правила оценки не изменяются ради улучшения результата отдельного инструмента.",
        "Кэш исключает необходимость повторных облачных вызовов при подготовке отчётов и TXT-экспортов.",
        "Web-приложение является демонстрационным; перед публичным размещением требуются аутентификация, ограничения запросов, постоянное хранилище заданий, полные lock-файлы и безопасное управление секретами.",
    ]:
        bullet(document, item)

    heading(document, "8. Основные папки проекта")
    add_table(document, ["Путь", "Содержимое"], [
        ["src/pdf_benchmark/", "Основной код адаптеров, benchmark, метрик и web-приложения."],
        ["documents/", "Пять исходных PDF тестового корпуса."],
        ["ground_truth/", "Эталонные объекты и их свойства для оценки."],
        ["outputs/benchmark/", "Версионированный кэш официальных экспериментов."],
        ["reports/", "Контрольные материалы и отчёты Prompt 12–20."],
        ["exports/txt/", "50 готовых TXT-файлов: по одному для каждой пары «конвертер–PDF»."],
    ])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("PDF Converter Benchmark · краткое описание архитектуры")
    set_font(run, 9)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.core_properties.title = "Краткое описание архитектуры проекта PDF Converter Benchmark"
    document.core_properties.author = "PDF Converter Benchmark"
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
