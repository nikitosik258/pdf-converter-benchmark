"""Create a two-section Word appendix for the final technical report."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "final_report" / "20261004T095716Z_s42_2334c932" / "architecture_appendix.docx"


def apply_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def add_paragraph(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.space_after = Pt(6)
    apply_font(paragraph.add_run(text), 12)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(6)
    apply_font(paragraph.add_run(text), 14, bold=True)


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_font(title.add_run("Дополнение к техническому отчёту"), 16, bold=True)

    add_heading(document, "1. Краткое описание архитектуры проекта")
    add_paragraph(
        document,
        "Проект построен как воспроизводимая система сравнения средств преобразования PDF. "
        "Исходные документы и эталонная разметка Ground Truth образуют неизменяемый тестовый корпус. "
        "Для каждого из десяти инструментов реализован адаптер с единым интерфейсом: он получает PDF и приводит ответ "
        "локальной библиотеки или облачного сервиса к общей структуре StandardizedDocument. В этой структуре хранятся "
        "текст, таблицы, математические и химические формулы, изображения, диаграммы и координаты объектов на страницах."
    )
    add_paragraph(
        document,
        "Модуль benchmark сопоставляет стандартизированные результаты с Ground Truth и рассчитывает оценки по шести категориям. "
        "Версионированный кэш хранит результаты обработки и позволяет повторять анализ без повторного запуска тяжёлых моделей "
        "и облачных API. Отдельный модуль формирует таблицы, графики, error analysis и итоговый отчёт. Web-приложение на FastAPI "
        "повторно использует те же адаптеры: пользователь загружает PDF, выбирает converter, получает статус задания и может скачать TXT или JSON."
    )

    add_heading(document, "2. Что и как делается в проекте")
    add_paragraph(
        document,
        "В исследовании используются пять технических PDF объёмом 98 страниц. В них заранее выделены 40 эталонных объектов: "
        "фрагменты обычного текста, таблицы, математические и химические формулы, изображения и диаграммы. Каждый PDF обрабатывается "
        "каждым из десяти инструментов: пятью локальными библиотеками и пятью облачными сервисами."
    )
    add_paragraph(
        document,
        "После обработки ответы приводятся к общему формату и автоматически сопоставляются с эталонными объектами по категории, "
        "содержимому и расположению на странице. На этой основе рассчитываются количественные показатели качества, а затем строятся "
        "сравнения по инструментам, типам контента и документам. В отчёте анализируются не только сводные баллы, но и специализация "
        "каждого инструмента, характерные ошибки, скорость, стоимость и требования к ресурсам. Для повторяемости все результаты, "
        "промежуточные данные и сформированные TXT-файлы сохраняются в проекте."
    )

    document.core_properties.title = "Дополнение: архитектура и процесс работы PDF Converter Benchmark"
    document.core_properties.author = "PDF Converter Benchmark"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
